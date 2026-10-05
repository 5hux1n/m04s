"""Generate anonymous fixtures inside the project; no printer or network."""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'console/vendor'))


def make_fixtures():
    import pymupdf
    from docx import Document
    from PIL import Image, ImageDraw

    directory = ROOT / 'artifacts/console/fixtures'
    directory.mkdir(parents=True, exist_ok=True)
    pdf = pymupdf.open()
    for number in (1, 2):
        page = pdf.new_page(width=300, height=400)
        page.insert_text((25, 40), f'M04S test document - page {number}', fontsize=14)
        page.insert_text((25, 80), 'TEST 12345678', fontsize=16)
    pdf.set_metadata({'title': 'M04S test document', 'author': 'm04s tests'})
    pdf.save(directory / '示例两页.pdf')
    pdf.close()
    document = Document()
    document.core_properties.author = 'm04s tests'
    document.core_properties.last_modified_by = 'm04s tests'
    document.add_heading('示例文档 TEST', 0)
    document.add_paragraph('这是用于导入验证的公开示例文字。12345678')
    document.add_picture(str(ROOT / 'samples/66.png'))
    document.save(directory / '示例文档.docx')
    (directory / '示例文字.txt').write_text('示例文字 TEST 12345678\n第二行', encoding='utf-8')
    image = Image.new('RGB', (240, 160))
    image.putdata([(x % 256, y * 255 // 159, (x + y) % 256) for y in range(160) for x in range(240)])
    draw = ImageDraw.Draw(image)
    draw.ellipse((30, 25, 130, 125), fill='white', outline='black', width=4)
    draw.text((45, 65), '66 TEST', fill='black')
    image.save(directory / '效果测试.png')
    return directory


if __name__ == '__main__':
    print(make_fixtures())
